"""
Роут: загрузка .docx с расшифровкой телефонного звонка, анализ через GenAPI
(gemini-3-5-flash-lite) и сохранение результата в таблицы call_summaries и examples.

Если для звонка уже существует сводка, старая сводка и её примеры удаляются,
после чего создаётся новая версия сводки.
"""

import io
import os
from pathlib import Path

import docx
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.crud.examples_crud import create_example
from app.crud.summaries_crud import create_summary, get_summary_by_call_id
from app.database import get_db
from app.models import Call, Example, ExampleCategory
from app.schemas import (
    CallSummaryCreate,
    CallSummaryDetailOut,
    CategoryOut,
    ExampleCreate,
)
from app.services.genapi_client import analyze_call_transcript


router = APIRouter(prefix="/calls", tags=["call-summary"])

FILES_DIR = Path(os.getenv("FILES_DIR", "files"))


def _extract_text_from_docx(file_bytes: bytes) -> str:
    """Извлекает весь текст (абзацы + таблицы) из .docx файла."""
    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Не удалось прочитать .docx файл: {exc}",
        )

    parts: list[str] = []

    # Текст обычных абзацев
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # Текст из таблиц
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    text = "\n".join(parts).strip()

    if not text:
        raise HTTPException(
            status_code=400,
            detail="В .docx файле не найден текстовый контент разговора.",
        )

    return text


@router.post(
    "/{call_id}/summary/upload-docx",
    response_model=CallSummaryDetailOut,
    status_code=status.HTTP_200_OK,
    summary="Загрузить .docx с расшифровкой звонка и получить сводку через GenAPI",
)
async def upload_call_docx_and_summarize(
    call_id: int,
    file: UploadFile = File(
        ...,
        description="Файл .docx с расшифровкой телефонного звонка",
    ),
    db: Session = Depends(get_db),
) -> CallSummaryDetailOut:
    """
    1. Принимает .docx с расшифровкой звонка для указанного call_id.
    2. Извлекает из него текст.
    3. Отправляет текст в GenAPI на анализ.
    4. Валидирует полученный JSON.
    5. Если сводка уже существует — удаляет старую сводку и её примеры.
    6. Сохраняет новую сводку.
    7. Сохраняет новые примеры.
    8. Возвращает новую сводку клиенту.
    """

    # ---------------------------------------------------------
    # 1. Проверяем расширение файла
    # ---------------------------------------------------------

    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(
            status_code=400,
            detail="Ожидается файл с расширением .docx",
        )

    # ---------------------------------------------------------
    # 2. Проверяем существование звонка
    # ---------------------------------------------------------

    call = db.query(Call).filter(Call.id == call_id).first()

    if call is None:
        raise HTTPException(
            status_code=404,
            detail="Звонок с указанным call_id не найден",
        )

    # ---------------------------------------------------------
    # 3. Читаем файл
    # ---------------------------------------------------------

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="Файл пустой",
        )

    # ---------------------------------------------------------
    # 4. Извлекаем текст из DOCX
    # ---------------------------------------------------------

    transcript_text = _extract_text_from_docx(file_bytes)

    # ---------------------------------------------------------
    # 5. Отправляем расшифровку в GenAPI
    # ---------------------------------------------------------

    analysis, error = await analyze_call_transcript(
        transcript_text=transcript_text,
        id_call=call_id,
    )

    if analysis is None:
        raise HTTPException(
            status_code=502,
            detail=f"Ошибка анализа GenAPI: {error}",
        )

    # ---------------------------------------------------------
    # 6. Удаляем старую сводку, если она существует
    # ---------------------------------------------------------

    existing_summary = get_summary_by_call_id(
        db,
        call_id,
    )

    if existing_summary is not None:

        # Сначала удаляем примеры старой сводки.
        db.execute(
            delete(Example).where(
                Example.summary_id == existing_summary.id
            )
        )

        # Затем удаляем саму старую сводку.
        db.delete(existing_summary)

        db.commit()

    # ---------------------------------------------------------
    # 7. Создаём новую сводку
    # ---------------------------------------------------------

    db_summary = create_summary(
        db,
        CallSummaryCreate(
            call_id=call_id,
            discussion=analysis["discussion"]["value"],
            agreements=analysis["agreements"]["value"],
            risks=analysis["risks"]["value"],
        ),
    )

    # ---------------------------------------------------------
    # 8. Сохраняем примеры новой сводки
    # ---------------------------------------------------------

    for category in (
        ExampleCategory.discussion,
        ExampleCategory.agreements,
        ExampleCategory.risks,
    ):
        category_examples = analysis[category.value].get(
            "examples",
            [],
        )

        for example_text in category_examples:

            if example_text and example_text.strip():

                create_example(
                    db,
                    ExampleCreate(
                        summary_id=db_summary.id,
                        category=category,
                        text=example_text,
                    ),
                )

    # ---------------------------------------------------------
    # 9. Обновляем объект сводки
    # ---------------------------------------------------------

    db.refresh(db_summary)

    # ---------------------------------------------------------
    # 10. Возвращаем результат в формате frontend
    # ---------------------------------------------------------

    return CallSummaryDetailOut(
        id=db_summary.id,
        call_id=call_id,
        discussion=CategoryOut(
            **analysis["discussion"],
        ),
        agreements=CategoryOut(
            **analysis["agreements"],
        ),
        risks=CategoryOut(
            **analysis["risks"],
        ),
    )

@router.post(
    "/{call_id}/summary/rebuild",
    response_model=CallSummaryDetailOut,
    status_code=status.HTTP_200_OK,
    summary="Пересобрать сводку звонка",
)
async def rebuild_call_summary(
    call_id: int,
    db: Session = Depends(get_db),
) -> CallSummaryDetailOut:
    """
    Пересобрать сводку существующего звонка.

    1. Находит звонок по call_id.
    2. Берёт путь к расшифровке из call.call_link.
    3. Читает существующий .docx.
    4. Извлекает текст расшифровки.
    5. Отправляет текст в GenAPI.
    6. Удаляет старую сводку и её примеры.
    7. Создаёт новую сводку.
    8. Сохраняет новые примеры.
    9. Возвращает новую сводку.
    """

    # 1. Проверяем существование звонка
    call = db.query(Call).filter(Call.id == call_id).first()

    if call is None:
        raise HTTPException(
            status_code=404,
            detail="Звонок с указанным call_id не найден",
        )

    # 2. Проверяем наличие ссылки на расшифровку
    if not call.call_link:
        raise HTTPException(
            status_code=404,
            detail="Для этого звонка не указан файл расшифровки",
        )

    # 3. Формируем путь к файлу
    #
    # call.call_link хранится, например:
    # files/2026_09_20_11_15_ivanov.docx
    #
    # В Docker FILES_DIR=/app/files.
    file_path = Path(call.call_link)

    if file_path.is_absolute():
        resolved_file_path = file_path
    elif file_path.parts and file_path.parts[0] == "files":
        resolved_file_path = FILES_DIR / file_path.relative_to("files")
    else:
        resolved_file_path = FILES_DIR / file_path

    file_path = resolved_file_path

    # 4. Проверяем существование файла
    if not file_path.exists():
        raise HTTPException(
            status_code=404,
            detail=f"Файл расшифровки не найден: {call.call_link}",
        )

    if not file_path.is_file():
        raise HTTPException(
            status_code=400,
            detail=f"Указанный путь не является файлом: {call.call_link}",
        )

    # 5. Проверяем расширение
    if file_path.suffix.lower() != ".docx":
        raise HTTPException(
            status_code=400,
            detail="Файл расшифровки должен иметь расширение .docx",
        )

    # 6. Читаем файл
    try:
        file_bytes = file_path.read_bytes()
    except OSError as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Не удалось прочитать файл расшифровки: {exc}",
        )

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="Файл расшифровки пустой",
        )

    # 7. Извлекаем текст из DOCX
    transcript_text = _extract_text_from_docx(file_bytes)

    # 8. Отправляем расшифровку в GenAPI
    analysis, error = await analyze_call_transcript(
        transcript_text=transcript_text,
        id_call=call_id,
    )

    if analysis is None:
        raise HTTPException(
            status_code=502,
            detail=f"Ошибка анализа GenAPI: {error}",
        )

    # 9. Находим существующую сводку
    existing_summary = get_summary_by_call_id(
        db,
        call_id,
    )

    # 10. Удаляем старую сводку и её примеры
    if existing_summary is not None:

        db.execute(
            delete(Example).where(
                Example.summary_id == existing_summary.id
            )
        )

        db.delete(existing_summary)

        db.commit()

    # 11. Создаём новую сводку
    db_summary = create_summary(
        db,
        CallSummaryCreate(
            call_id=call_id,
            discussion=analysis["discussion"]["value"],
            agreements=analysis["agreements"]["value"],
            risks=analysis["risks"]["value"],
        ),
    )

    # 12. Сохраняем примеры новой сводки
    for category in (
        ExampleCategory.discussion,
        ExampleCategory.agreements,
        ExampleCategory.risks,
    ):
        category_examples = analysis[category.value].get(
            "examples",
            [],
        )

        for example_text in category_examples:

            if example_text and example_text.strip():

                create_example(
                    db,
                    ExampleCreate(
                        summary_id=db_summary.id,
                        category=category,
                        text=example_text,
                    ),
                )

    # 13. Обновляем объект из БД
    db.refresh(db_summary)

    # 14. Возвращаем новую сводку
    return CallSummaryDetailOut(
        id=db_summary.id,
        call_id=call_id,
        discussion=CategoryOut(
            **analysis["discussion"],
        ),
        agreements=CategoryOut(
            **analysis["agreements"],
        ),
        risks=CategoryOut(
            **analysis["risks"],
        ),
    )